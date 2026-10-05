"""Parser de índices XLSX del expediente judicial.

El XLSX es la fuente de verdad procesal (plan §1.1, decisión D5). Este módulo
solo lee y normaliza; NO empareja con archivos físicos (eso hace el importador).
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterator

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet


@dataclass(frozen=True, slots=True)
class IndexEntry:
    """Una fila del índice que describe un ítem procesal."""

    indice_numero: int | None
    nombre_original: str
    fecha_creacion: date | None
    fecha_incorporacion: date | None
    orden_documento: int | None
    numero_paginas: int | None
    pagina_inicio: int | None
    pagina_fin: int | None
    formato: str | None


@dataclass(frozen=True, slots=True)
class IndexMetadata:
    """Metadatos de la carátula del índice."""

    radicacion: str | None
    ciudad: str | None
    despacho: str | None
    serie: str | None
    parte_a: str | None
    parte_b: str | None


@dataclass(frozen=True, slots=True)
class CuadernoIndex:
    """Índice completo de un cuaderno."""

    cuaderno: str
    metadata: IndexMetadata
    entries: list[IndexEntry]


@dataclass(frozen=True, slots=True)
class GeneralIndex:
    """Índice maestro general: lista de cuadernos del expediente."""

    metadata: IndexMetadata
    cuadernos: list["CuadernoRef"]


@dataclass(frozen=True, slots=True)
class CuadernoRef:
    """Referencia a un cuaderno en el índice general."""

    orden: int | None
    nombre: str


# Filas fijas observadas en los XLSX reales (1-based).
_CARATULA_ROWS = {
    "ciudad": 2,
    "despacho": 3,
    "serie": 4,
    "radicacion": 5,
    "parte_a": 6,
    "parte_b": 7,
}
_HEADER_ROW = 9
_DATA_START = 10


def _normalize_text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def _to_int(value: object) -> int | None:
    if value is None:
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value == int(value) else None
    m = re.search(r"\d+", str(value))
    return int(m.group()) if m else None


def _to_date(value: object) -> date | None:
    """Acepta datetime/date de Excel o cadenas con día/mes/año."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    # Limpia cadenas como "06//04/2022" o "06/04/2022"
    cleaned = re.sub(r"[^0-9]", " ", text).split()
    if len(cleaned) >= 3:
        try:
            day, month, year = int(cleaned[0]), int(cleaned[1]), int(cleaned[2])
            if year < 100:
                year += 2000
            return date(year, month, day)
        except ValueError:
            return None
    return None


def _extract_index_number(name: str) -> int | None:
    """Extrae el número inicial de un nombre de índice (ej: '0003Auto...' -> 3)."""
    m = re.match(r"^(\d+)", name)
    return int(m.group(1)) if m else None


def _caratula_value(ws: Worksheet, row: int) -> str | None:
    # Columna B (2) es el valor; columna A es la etiqueta
    return _normalize_text(ws.cell(row=row, column=2).value)


def _parse_metadata(ws: Worksheet) -> IndexMetadata:
    return IndexMetadata(
        ciudad=_caratula_value(ws, _CARATULA_ROWS["ciudad"]),
        despacho=_caratula_value(ws, _CARATULA_ROWS["despacho"]),
        serie=_caratula_value(ws, _CARATULA_ROWS["serie"]),
        radicacion=_caratula_value(ws, _CARATULA_ROWS["radicacion"]),
        parte_a=_caratula_value(ws, _CARATULA_ROWS["parte_a"]),
        parte_b=_caratula_value(ws, _CARATULA_ROWS["parte_b"]),
    )


def _parse_row(ws: Worksheet, row: int) -> IndexEntry | None:
    name = _normalize_text(ws.cell(row=row, column=1).value)
    # Ignora la fila de cierre y filas vacías
    if not name or name.upper().startswith("FECHA DE CIERRE"):
        return None
    fmt = _normalize_text(ws.cell(row=row, column=8).value)
    orden = _to_int(ws.cell(row=row, column=4).value)
    indice_numero = _extract_index_number(name) or orden
    return IndexEntry(
        indice_numero=indice_numero,
        nombre_original=name,
        fecha_creacion=_to_date(ws.cell(row=row, column=2).value),
        fecha_incorporacion=_to_date(ws.cell(row=row, column=3).value),
        orden_documento=orden,
        numero_paginas=_to_int(ws.cell(row=row, column=5).value),
        pagina_inicio=_to_int(ws.cell(row=row, column=6).value),
        pagina_fin=_to_int(ws.cell(row=row, column=7).value),
        formato=fmt.lower() if fmt else None,
    )


def parse_cuaderno_index(path: Path, cuaderno: str) -> CuadernoIndex:
    """Parsea un `0000IndiceExpedienteElectronico.xlsx` de un cuaderno."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    entries: list[IndexEntry] = []
    for row in range(_DATA_START, ws.max_row + 1):
        entry = _parse_row(ws, row)
        if entry is not None:
            entries.append(entry)
    return CuadernoIndex(cuaderno=cuaderno, metadata=_parse_metadata(ws), entries=entries)


def parse_general_index(path: Path) -> GeneralIndex:
    """Parsea el `0000IndiceExpedienteGeneral.xlsx` de la raíz del expediente."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    cuadernos: list[CuadernoRef] = []
    for row in range(_DATA_START, ws.max_row + 1):
        name = _normalize_text(ws.cell(row=row, column=1).value)
        if not name or name.upper().startswith("FECHA DE CIERRE"):
            continue
        orden = _to_int(ws.cell(row=row, column=4).value)
        cuadernos.append(CuadernoRef(orden=orden, nombre=name))
    return GeneralIndex(metadata=_parse_metadata(ws), cuadernos=cuadernos)


def iter_expediente_indices(root: Path) -> Iterator[CuadernoIndex]:
    """Recorre un expediente y produce el índice de cada cuaderno.

    Estructura esperada:
        root/
          0000IndiceExpedienteGeneral.xlsx
          01PrimeraInstancia/
            <carpeta cuaderno>/
              0000IndiceExpedienteElectronico.xlsx
              ...
    """
    for instancia in sorted(root.iterdir()):
        if not instancia.is_dir():
            continue
        for cuaderno_dir in sorted(instancia.iterdir()):
            if not cuaderno_dir.is_dir():
                continue
            index_file = cuaderno_dir / "0000IndiceExpedienteElectronico.xlsx"
            if index_file.exists():
                cuaderno_name = f"{instancia.name}/{cuaderno_dir.name}"
                yield parse_cuaderno_index(index_file, cuaderno_name)
