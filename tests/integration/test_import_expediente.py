"""IT-IMP — importador masivo de expediente (Fase 1)."""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from tests.helpers import create_case, new_case_number

pytestmark = pytest.mark.integration


def _make_pdf(path: Path, text: str = "demanda") -> None:
    from seeds.pdfgen import make_pdf
    path.write_bytes(make_pdf([text]))


def _make_mp4(path: Path) -> None:
    # Firma mínima MP4: ftyp + tamaño + brand
    path.write_bytes(b"\x00\x00\x00\x20ftypisom\x00\x00\x00\x00isommp41")


def _make_xlsx_index(path: Path, entries: list[tuple[int, str, str]], radicacion: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Índice Electrónico"
    ws["A5"] = "No. Radicación  del Proceso"
    ws["B5"] = radicacion
    ws["A9"] = "Nombre del Documento"
    ws["D9"] = "Orden Documento"
    ws["H9"] = "Formato"
    for i, (orden, name, fmt) in enumerate(entries, start=10):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=4, value=orden)
        ws.cell(row=i, column=8, value=fmt)
    wb.save(path)


def _build_mini_expediente(root: Path, radicacion: str) -> None:
    root.mkdir(parents=True)
    # Índice general
    general = root / "0000IndiceExpedienteGeneral.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A5"] = "No. Radicación"
    ws["B5"] = radicacion
    ws["A9"] = "Nombre"
    ws["D9"] = "Orden"
    ws["A10"] = "01PrimeraInstancia"
    ws["D10"] = 1
    wb.save(general)

    # Cuaderno 0001 con 2 PDFs
    c1 = root / "01PrimeraInstancia" / "0001 DemandaPrincipal2018-361"
    c1.mkdir(parents=True)
    _make_xlsx_index(c1 / "0000IndiceExpedienteElectronico.xlsx",
                     [(1, "0001 DemandaPrincipal2018-361", "pdf"),
                      (2, "0002AutoAdmite", "pdf")],
                     radicacion)
    _make_pdf(c1 / "0001 DemandaPrincipal2018-361.pdf", "demanda principal")
    _make_pdf(c1 / "0002AutoAdmite.pdf", "auto admite")

    # Cuaderno 0002 con subcarpeta de diligencia (múltiples archivos, 1 video)
    c2 = root / "01PrimeraInstancia" / "0002 MedidasCautelares2018-361"
    c2.mkdir(parents=True)
    _make_xlsx_index(c2 / "0000IndiceExpedienteElectronico.xlsx",
                     [(1, "0001 CuadernoMedidasCautelares2018-361", "pdf"),
                      (2, "0002DiligenciaComisorio", "pdf")],
                     radicacion)
    _make_pdf(c2 / "0001 CuadernoMedidasCautelares2018-361.pdf", "medidas")
    dil = c2 / "0002DiligenciaComisorio"
    dil.mkdir()
    _make_pdf(dil / "001Acta.pdf", "acta diligencia")
    _make_mp4(dil / "002Video.mp4")


def test_it_imp_01_imports_full_expediente(client, auth, tmp_path: Path, owner_db):
    h = auth("admin.alfa")
    radicacion = new_case_number()
    case = create_case(client, h, "Import test", case_number=radicacion)
    case_id = case["id"]
    org_id = client.get("/v1/auth/me", headers=h).json()["organization_id"]
    user_id = client.get("/v1/auth/me", headers=h).json()["id"]

    root = tmp_path / "expediente"
    _build_mini_expediente(root, radicacion)

    from scripts.import_expediente import import_expediente
    result = import_expediente(folder=root, case_id=case_id, org_id=org_id, user_id=user_id)

    assert result.registered == 8  # 2 XLSX índice + 1 general + 2 docs c1 + 1 doc c2 + 2 en subcarpeta
    assert result.duplicates == 0
    assert not result.errors

    # Verifica linaje en BD
    with owner_db.cursor() as cur:
        cur.execute("SELECT cuaderno, indice_numero, orden_procesal FROM documents WHERE case_id = %s ORDER BY orden_procesal",
                    (case_id,))
        docs = cur.fetchall()
        cur.execute("SELECT cuaderno, indice_numero, orden_procesal FROM media WHERE case_id = %s", (case_id,))
        media = cur.fetchall()
    owner_db.rollback()

    assert len(docs) == 7
    assert len(media) == 1
    # El índice general es maestro
    assert any(d[0] is None and d[1] is None for d in docs)
    # Los documentos del cuaderno 0001 tienen índices 1 y 2 (el XLSX del cuaderno es indice_numero=None)
    c1_docs = [d for d in docs if d[0] and d[0].endswith("0001 DemandaPrincipal2018-361")]
    assert {d[1] for d in c1_docs if d[1] is not None} == {1, 2}
    # Subcarpeta hereda índice 2
    c2_sub = [d for d in docs if d[0] and d[0].endswith("0002 MedidasCautelares2018-361") and d[1] == 2]
    assert len(c2_sub) == 1  # el PDF; el video va a media


def test_it_imp_02_dry_run_does_not_write(client, auth, tmp_path: Path, owner_db):
    h = auth("admin.alfa")
    radicacion = new_case_number()
    case = create_case(client, h, "Dry run test", case_number=radicacion)
    case_id = case["id"]
    org_id = client.get("/v1/auth/me", headers=h).json()["organization_id"]
    user_id = client.get("/v1/auth/me", headers=h).json()["id"]

    root = tmp_path / "expediente"
    _build_mini_expediente(root, radicacion)

    from scripts.import_expediente import import_expediente
    result = import_expediente(folder=root, case_id=case_id, org_id=org_id, user_id=user_id, dry_run=True)
    assert result.registered == 8

    with owner_db.cursor() as cur:
        cur.execute("SELECT count(*) FROM documents WHERE case_id = %s", (case_id,))
        assert cur.fetchone()[0] == 0
        cur.execute("SELECT count(*) FROM media WHERE case_id = %s", (case_id,))
        assert cur.fetchone()[0] == 0
    owner_db.rollback()


def test_it_imp_03_deduplicates_by_sha256(client, auth, tmp_path: Path, owner_db):
    h = auth("admin.alfa")
    radicacion = new_case_number()
    case = create_case(client, h, "Dedup test", case_number=radicacion)
    case_id = case["id"]
    org_id = client.get("/v1/auth/me", headers=h).json()["organization_id"]
    user_id = client.get("/v1/auth/me", headers=h).json()["id"]

    root = tmp_path / "expediente"
    _build_mini_expediente(root, radicacion)
    # Copia exacta en otro cuaderno para forzar dedup
    dup_dir = root / "01PrimeraInstancia" / "0003 Duplicado"
    dup_dir.mkdir(parents=True)
    (root / "01PrimeraInstancia" / "0001 DemandaPrincipal2018-361" / "0001 DemandaPrincipal2018-361.pdf").copy(
        dup_dir / "0001MismoDoc.pdf")

    from scripts.import_expediente import import_expediente
    result = import_expediente(folder=root, case_id=case_id, org_id=org_id, user_id=user_id)
    assert result.registered == 8  # 8 originales; el duplicado se detecta y no se registra
    assert result.duplicates == 1
