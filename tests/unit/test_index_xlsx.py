"""UT-IDX — parser de índices XLSX (Fase 1)."""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

from app.services.index_xlsx import (
    CuadernoRef,
    parse_cuaderno_index,
    parse_general_index,
)

pytestmark = pytest.mark.unit


def _make_cuaderno_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Índice Electrónico"
    ws["A2"] = "Ciudad"
    ws["B2"] = "BOGOTÁ"
    ws["A3"] = "Despacho Judicial"
    ws["B3"] = "JUZGADO 21"
    ws["A4"] = "Serie"
    ws["B4"] = "CIVIL"
    ws["A5"] = "No. Radicación"
    ws["B5"] = "11001310302120180036100"
    ws["A6"] = "Parte A"
    ws["B6"] = "Demandado"
    ws["A7"] = "Parte B"
    ws["B7"] = "Demandante"
    ws["A9"] = "Nombre del Documento"
    ws["B9"] = "Fecha Creación Documento"
    ws["C9"] = "Fecha Incorporación Expediente"
    ws["D9"] = "Orden Documento"
    ws["E9"] = "Número Paginas"
    ws["F9"] = "Página Inicio"
    ws["G9"] = "Página Fin"
    ws["H9"] = "Formato"
    ws["A10"] = "0001 DemandaPrincipal"
    ws["B10"] = datetime(2018, 7, 19)
    ws["D10"] = 1
    ws["E10"] = 10
    ws["H10"] = "pdf"
    ws["A11"] = "0002AutoInforma"
    ws["B11"] = "06//04/2022"
    ws["D11"] = 2
    ws["H11"] = "PDF"
    ws["A12"] = "0003CorreoInforma"
    ws["D12"] = 3
    ws["H12"] = "pdf"
    ws["A13"] = "0003 CorreoExtra"
    ws["D13"] = 3
    ws["H13"] = "pdf"  # duplicado de orden
    ws["A242"] = "FECHA DE CIERRE DEL EXPEDIENTE"
    wb.save(path)


def _make_general_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Índice Electrónico"
    ws["A2"] = "Ciudad"
    ws["B2"] = "BOGOTÁ"
    ws["A5"] = "No. Radicación"
    ws["B5"] = "11001310302120180036100"
    ws["A9"] = "Nombre del Documento"
    ws["D9"] = "Orden Documento"
    ws["A10"] = "01PrimeraInstancia"
    ws["D10"] = 1
    ws["A11"] = "02SegundaInstancia"
    ws["D11"] = 2
    ws["A242"] = "FECHA DE CIERRE DEL EXPEDIENTE"
    wb.save(path)


def test_ut_idx_01_parse_cuaderno_excludes_empty_and_closing(tmp_path: Path):
    p = tmp_path / "0000IndiceExpedienteElectronico.xlsx"
    _make_cuaderno_xlsx(p)
    idx = parse_cuaderno_index(p, "01PrimeraInstancia/0001 X")
    assert idx.cuaderno == "01PrimeraInstancia/0001 X"
    assert len(idx.entries) == 4
    assert idx.metadata.radicacion == "11001310302120180036100"


def test_ut_idx_02_extracts_numbers_and_lowercases_format(tmp_path: Path):
    p = tmp_path / "0000IndiceExpedienteElectronico.xlsx"
    _make_cuaderno_xlsx(p)
    idx = parse_cuaderno_index(p, "c")
    e1, e2 = idx.entries[0], idx.entries[1]
    assert e1.indice_numero == 1
    assert e1.nombre_original == "0001 DemandaPrincipal"
    assert e1.formato == "pdf"
    assert e1.fecha_creacion == date(2018, 7, 19)
    assert e2.indice_numero == 2
    assert e2.formato == "pdf"  # normalizado a lowercase
    assert e2.fecha_creacion == date(2022, 4, 6)  # "06//04/2022"


def test_ut_idx_03_keeps_duplicate_orders_as_separate_entries(tmp_path: Path):
    p = tmp_path / "0000IndiceExpedienteElectronico.xlsx"
    _make_cuaderno_xlsx(p)
    idx = parse_cuaderno_index(p, "c")
    ords = [e.indice_numero for e in idx.entries]
    assert ords == [1, 2, 3, 3]


def test_ut_idx_04_parse_general_lists_cuadernos(tmp_path: Path):
    p = tmp_path / "0000IndiceExpedienteGeneral.xlsx"
    _make_general_xlsx(p)
    g = parse_general_index(p)
    assert len(g.cuadernos) == 2
    assert g.cuadernos[0] == CuadernoRef(orden=1, nombre="01PrimeraInstancia")
    assert g.cuadernos[1] == CuadernoRef(orden=2, nombre="02SegundaInstancia")
    assert g.metadata.radicacion == "11001310302120180036100"
